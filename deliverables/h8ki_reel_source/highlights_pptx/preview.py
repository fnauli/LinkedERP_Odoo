"""Rough PNG preview of the deck (LibreOffice can't run here): shapes, pictures and wrapped text in
Liberation Sans (Arial metrics), so positions, overflow and wrapping can be checked. Canvas includes the
area right of the slide, where the sticky notes sit."""
import sys, io, math
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.dml import MSO_THEME_COLOR, MSO_FILL
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

THEME = {"ACCENT_1": "E3181E", "ACCENT_2": "049844", "ACCENT_3": "291B73", "ACCENT_4": "FFB400",
         "ACCENT_5": "966E00", "ACCENT_6": "FFF3B0", "TEXT_1": "181C28", "DARK_1": "181C28",
         "TEXT_2": "606874", "DARK_2": "606874", "BACKGROUND_1": "FFFFFF", "LIGHT_1": "FFFFFF",
         "BACKGROUND_2": "F6F7FA", "LIGHT_2": "F6F7FA"}
FD = "/usr/share/fonts/truetype/liberation/LiberationSans-"
PX = 96 / 914400


def F(sz, b, i=False):
    return ImageFont.truetype(FD + ("BoldItalic" if b and i else "Bold" if b else "Italic" if i else "Regular") + ".ttf",
                              max(1, int(round(sz))))


def col(c, default="000000"):
    try:
        if c.type is None:
            return default
        if c.theme_color is not None and c.type == 2:
            return THEME[MSO_THEME_COLOR(c.theme_color).name]
        return str(c.rgb)
    except Exception:
        return default


def rgb(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def bg_of(slide):
    f = slide.background.fill
    try:
        return rgb(col(f.fore_color, "F6F7FA"))
    except Exception:
        return rgb("F6F7FA")


def draw(slide, W=1080, H=1920, extra=520):
    im = Image.new("RGB", (W + extra, H), (200, 200, 200))
    im.paste(bg_of(slide), (0, 0, W, H))
    d = ImageDraw.Draw(im)
    issues = []
    for sh in slide.shapes:
        x, y, w, h = [int(round(v * PX)) for v in (sh.left, sh.top, sh.width, sh.height)]
        if sh.shape_type == 13:
            pic = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA").resize((w, h), Image.LANCZOS)
            im.paste(pic, (x, y), pic)
            continue
        geom = None
        try:
            geom = sh.auto_shape_type
        except Exception:
            pass
        fill = None
        try:
            if sh.fill.type == MSO_FILL.SOLID:
                fill = rgb(col(sh.fill.fore_color))
        except Exception:
            pass
        name = str(geom) if geom is not None else ""
        if fill:
            if "ROUNDED" in name:
                adj = sh.adjustments[0] if len(sh.adjustments) else 0.16667
                d.rounded_rectangle((x, y, x + w, y + h), int(adj * min(w, h)), fill=fill)
            elif "OVAL" in name:
                d.ellipse((x, y, x + w, y + h), fill=fill)
            elif "STAR" in name:
                cx, cy = x + w / 2, y + h / 2
                d.polygon([(cx + (w / 2 if k % 2 == 0 else w * 0.19) * math.cos(-math.pi / 2 + k * math.pi / 5),
                            cy + (h / 2 if k % 2 == 0 else h * 0.19) * math.sin(-math.pi / 2 + k * math.pi / 5))
                           for k in range(10)], fill=fill)
            else:
                d.rectangle((x, y, x + w, y + h), fill=fill)
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        tf = sh.text_frame
        ml, mr, mt, mb = [int((v or 0) * PX) for v in (tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom)]
        bx, bw = x + ml, w - ml - mr
        lines = []
        for p in tf.paragraphs:
            runs = [(r.text, r.font) for r in p.runs]
            if not runs:
                continue
            f0 = runs[0][1]
            sz = (f0.size.pt if f0.size else 18) / 0.75
            fnt = F(sz, bool(f0.bold), bool(f0.italic))
            color = rgb(col(f0.color, "181C28"))
            text = "".join(t for t, _ in runs)
            bullet = p._p.pPr is not None and p._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
            indent = int(sz * 1.2) if bullet else 0
            cur = ""
            wl = []
            for word in text.split(" "):
                t = (cur + " " + word) if cur else word
                if fnt.getlength(t) <= bw - indent:
                    cur = t
                else:
                    if cur:
                        wl.append(cur)
                    cur = word
                    if fnt.getlength(word) > bw - indent:
                        issues.append(f"word too wide: {word!r}")
            wl.append(cur)
            sb = (p.space_before.pt / 0.75) if p.space_before else 0
            sa = (p.space_after.pt / 0.75) if p.space_after else 0
            for k, l in enumerate(wl):
                lines.append((("• " if bullet and k == 0 else "") + l, fnt, color, sz * 1.2, sb if k == 0 else 0,
                              p.alignment, indent if k else 0))
            if sa:
                lines.append(("", fnt, color, 0, sa, None, 0))
        total = sum(lh + sb for _, _, _, lh, sb, _, _ in lines)
        avail = h - mt - mb
        if total > avail + 2:
            issues.append(f"overflow {sh.name!r}: text {total:.0f}px > box {avail}px")
        va = tf.vertical_anchor
        yy = y + mt + (avail - total) / 2 if va == MSO_ANCHOR.MIDDLE else y + mt
        for l, fnt, color, lh, sb, al, ind in lines:
            yy += sb
            if l:
                tw = fnt.getlength(l)
                xx = bx + (bw - tw) / 2 if al == PP_ALIGN.CENTER else bx + ind
                d.text((xx, yy + (lh - fnt.size) * 0.5), l, font=fnt, fill=color)
            yy += lh
        if x + w > W and x < W:
            issues.append(f"crosses slide edge: {sh.name!r}")
    d.line([(W, 0), (W, H)], fill=(0, 0, 0), width=3)
    return im, issues


prs = Presentation(sys.argv[1])
out = []
for i, s in enumerate(prs.slides, 1):
    im, iss = draw(s)
    im.save(f"prev_{i}.png")
    out.append(im)
    print(i, iss or "ok")
sheet = Image.new("RGB", (len(out) * 400, 480), (255, 255, 255))
for k, im in enumerate(out):
    sheet.paste(im.crop((0, 0, 1080, 1920)).resize((270, 480)), (k * 400, 0))
sheet.save("prev_sheet.png")
